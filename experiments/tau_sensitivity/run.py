"""τ sensitivity analysis (manuscript §3.2.6, §4.1).

Loads the 30-query benchmark register, sweeps τ around the recovered value from
calib_db, and writes results.csv plus figure3.pdf/.png.
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import config
from ara.complexity import DEFAULT_TAU

OUT_DIR = Path(__file__).resolve().parent
RESULTS_CSV = OUT_DIR / "results.csv"
FIG_PDF = OUT_DIR / "figure3.pdf"
FIG_PNG = OUT_DIR / "figure3.png"
PLAN_MD = OUT_DIR / "PLAN.md"

BENCHMARK_CANDIDATES = (
    ROOT / "data" / "ARA_benchmark_30queries.xlsx",
    ROOT / "data" / "ARA_Benchmark_30_Queries.xlsx",
)

REGISTER_COLUMNS = (
    "query_id",
    "design_stratum",
    "ground_truth_class",
    "computed_score",
)

N_TAU = 13


def resolve_benchmark_path() -> Path:
    for path in BENCHMARK_CANDIDATES:
        if path.is_file():
            return path
    names = ", ".join(p.name for p in BENCHMARK_CANDIDATES)
    raise FileNotFoundError(f"Benchmark workbook not found (tried {names})")


def read_recovered_tau(calib_db_dir: Path | None = None) -> float:
    """Load recovered τ from calibration artifacts tied to calib_db."""
    calib_db = calib_db_dir or Path(config.CALIB_DB_DIR)
    if not calib_db.is_dir():
        raise FileNotFoundError(f"calib_db directory not found: {calib_db}")

    candidates = [
        calib_db / "learned_weights.json",
        Path(config.LEARNED_WEIGHTS_PATH),
    ]
    seen: set[Path] = set()
    for path in candidates:
        path = path.resolve()
        if path in seen or not path.is_file():
            continue
        seen.add(path)
        data = json.loads(path.read_text(encoding="utf-8"))
        if "tau" in data:
            return float(data["tau"])

    for path in sorted(calib_db.glob("*.json")):
        if path.name == "corpus_stats.json":
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict) and "tau" in data:
            return float(data["tau"])

    raise FileNotFoundError(
        f"No JSON with 'tau' under {calib_db} or {config.LEARNED_WEIGHTS_PATH}"
    )


def read_tau_paper(wb_path: Path) -> float:
    """Paper-scale τ from the benchmark Model sheet (fallback: DEFAULT_TAU)."""
    wb = load_workbook(wb_path, read_only=True, data_only=True)
    try:
        if "Model" in wb.sheetnames:
            ws = wb["Model"]
            for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
                label, value = row[0], row[1]
                if label and "Routing threshold" in str(label) and value is not None:
                    return float(value)
    finally:
        wb.close()
    return DEFAULT_TAU


def _normalize_header(cell) -> str:
    return str(cell).replace("\n", " ").strip() if cell else ""


def _pick_column(headers: list[str], *candidates: str) -> str | None:
    lower = {h.lower(): h for h in headers}
    for name in candidates:
        if name.lower() in lower:
            return lower[name.lower()]
    return None


def load_query_register(wb_path: Path) -> pd.DataFrame:
    wb = load_workbook(wb_path, read_only=True, data_only=True)
    try:
        if "Query_Register" in wb.sheetnames:
            ws = wb["Query_Register"]
            rows = list(ws.iter_rows(values_only=True))
            header_row = next(
                i for i, row in enumerate(rows)
                if row and _normalize_header(row[0]).lower() == "query_id"
            )
            headers = [_normalize_header(h) for h in rows[header_row]]
            alias = {
                "query_id": _pick_column(headers, "query_id", "ID"),
                "design_stratum": _pick_column(
                    headers, "design_stratum", "Intended Stratum"
                ),
                "ground_truth_class": _pick_column(
                    headers, "ground_truth_class", "GT Class"
                ),
                "computed_score": _pick_column(
                    headers, "computed_score", "C(Q)", "Recal C(Q)"
                ),
            }
            if any(v is None for v in alias.values()):
                missing = [k for k, v in alias.items() if v is None]
                raise ValueError(f"Query_Register missing columns: {missing}")

            records = []
            for row in rows[header_row + 1 :]:
                if not row or row[headers.index(alias["query_id"])] is None:
                    continue
                rec = {
                    key: row[headers.index(col)]
                    for key, col in alias.items()
                }
                if rec["computed_score"] is None:
                    continue
                records.append(rec)
        elif "Benchmark_30" in wb.sheetnames:
            ws = wb["Benchmark_30"]
            rows = list(ws.iter_rows(values_only=True))
            headers = [_normalize_header(h) for h in rows[3]]
            ix = {name: i for i, name in enumerate(headers)}
            records = []
            for row in rows[4:]:
                if not row[ix["ID"]]:
                    continue
                records.append(
                    {
                        "query_id": row[ix["ID"]],
                        "design_stratum": row[ix["Intended Stratum"]],
                        "ground_truth_class": row[ix["GT Class"]],
                        "computed_score": row[ix["C(Q)"]],
                    }
                )
        else:
            raise ValueError(
                f"{wb_path.name} has no Query_Register or Benchmark_30 sheet"
            )
    finally:
        wb.close()

    df = pd.DataFrame(records)
    if len(df) != 30:
        raise ValueError(f"Expected 30 queries, got {len(df)}")
    df["computed_score"] = pd.to_numeric(df["computed_score"], errors="raise")
    df["ground_truth_class"] = df["ground_truth_class"].astype(str).str.strip().str.lower()
    df["design_stratum"] = df["design_stratum"].astype(str).str.strip().str.lower()
    return df


def align_scores_to_recovered_scale(
    scores: np.ndarray, tau_recovered: float, tau_paper: float
) -> np.ndarray:
    """Map paper-scale register scores to the recovered-τ scale when needed.

    Query_Register / Benchmark_30 ``computed_score`` values use the paper α weights;
    recovered τ from calibrate.py lives on the re-fit scale.  Rescaling by
    τ_recovered / τ_paper preserves routing decisions at the operating point.
    """
    scores = np.asarray(scores, dtype=float)
    if np.median(scores) < tau_recovered / 2:
        return scores * (tau_recovered / tau_paper)
    return scores


def predict_route(scores: np.ndarray, tau: float) -> np.ndarray:
    return np.where(scores > tau, "complex", "simple")


def routing_accuracy(
    scores: np.ndarray, ground_truth: np.ndarray, tau: float
) -> tuple[float, int]:
    predicted = predict_route(scores, tau)
    correct = int(np.sum(predicted == ground_truth))
    return correct / len(ground_truth), correct


def sweep_tau(
    scores: np.ndarray,
    ground_truth: np.ndarray,
    strata: np.ndarray,
    tau_recovered: float,
) -> pd.DataFrame:
    taus = np.linspace(0.7 * tau_recovered, 1.3 * tau_recovered, N_TAU)
    borderline = strata == "borderline"
    rows = []
    for tau in taus:
        acc_all, n_all = routing_accuracy(scores, ground_truth, tau)
        acc_bl, n_bl = routing_accuracy(
            scores[borderline], ground_truth[borderline], tau
        )
        rows.append(
            {
                "tau": round(float(tau), 4),
                "accuracy_overall": round(acc_all, 4),
                "accuracy_borderline": round(acc_bl, 4),
                "n_correct_overall": n_all,
                "n_correct_borderline": n_bl,
            }
        )
    return pd.DataFrame(rows)


def plot_sensitivity(
    results: pd.DataFrame, tau_recovered: float, out_pdf: Path, out_png: Path
) -> None:
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial"],
            "font.size": 10,
            "axes.labelsize": 10,
            "legend.fontsize": 10,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
        }
    )

    fig, ax = plt.subplots(figsize=(6, 4))
    ax.plot(
        results["tau"],
        results["accuracy_overall"],
        "o-",
        color="black",
        label="Overall",
    )
    ax.plot(
        results["tau"],
        results["accuracy_borderline"],
        "o--",
        color="0.45",
        label="Borderline",
    )
    ax.axvline(tau_recovered, color="0.35", linestyle="--", linewidth=1)
    y_top = results[["accuracy_overall", "accuracy_borderline"]].max().max()
    ax.annotate(
        f"recovered τ = {tau_recovered:.2f}",
        xy=(tau_recovered, y_top),
        xytext=(8, 0),
        textcoords="offset points",
        ha="left",
        va="center",
        fontsize=10,
    )

    ax.set_xlabel("τ")
    ax.set_ylabel("Routing accuracy")
    ax.set_ylim(0, 1)
    ax.set_xlim(results["tau"].min(), results["tau"].max())
    ax.grid(True, linestyle="-", linewidth=0.5, alpha=0.6)
    ax.legend(loc="best", frameon=False)
    fig.tight_layout()
    fig.savefig(out_pdf, bbox_inches="tight")
    fig.savefig(out_png, dpi=300, bbox_inches="tight")
    plt.close(fig)


def write_plan(tau_recovered: float, run_date: str) -> None:
    text = f"""# τ sensitivity run log

- **Run date:** {run_date}
- **Recovered τ (runtime):** {tau_recovered:.4f} (from `{config.CALIB_DB_DIR}` / `{config.LEARNED_WEIGHTS_PATH}`)
- **Random seed:** none (deterministic sweep over fixed τ grid)
- **NumPy version:** {np.__version__}
"""
    PLAN_MD.write_text(text, encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    benchmark_path = resolve_benchmark_path()
    tau_recovered = read_recovered_tau()
    tau_paper = read_tau_paper(benchmark_path)

    register = load_query_register(benchmark_path)
    scores = align_scores_to_recovered_scale(
        register["computed_score"].to_numpy(),
        tau_recovered,
        tau_paper,
    )
    ground_truth = register["ground_truth_class"].to_numpy()
    strata = register["design_stratum"].to_numpy()

    results = sweep_tau(scores, ground_truth, strata, tau_recovered)
    results.to_csv(RESULTS_CSV, index=False)

    plot_sensitivity(results, tau_recovered, FIG_PDF, FIG_PNG)
    write_plan(tau_recovered, date.today().isoformat())

    closest = results.iloc[(results["tau"] - tau_recovered).abs().argmin()]
    print(f"Benchmark: {benchmark_path}")
    print(f"Recovered tau = {tau_recovered:.2f} (paper tau = {tau_paper:.2f})")
    print(f"Wrote {RESULTS_CSV} ({len(results)} rows)")
    print(f"Wrote {FIG_PDF} and {FIG_PNG}")
    print(f"Wrote {PLAN_MD}")
    print(
        "At tau nearest recovered: "
        f"accuracy_overall={closest['accuracy_overall']:.2f}, "
        f"accuracy_borderline={closest['accuracy_borderline']:.2f}"
    )


if __name__ == "__main__":
    main()
